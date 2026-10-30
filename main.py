import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import os
import csv
import sys
import threading
import re
from openpyxl import load_workbook


def resource_path(relative_path):
    # UNC网络共享路径直接返回，不拼接本地目录
    if relative_path.startswith(r"\\"):
        return relative_path
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)


# ===== 导入配置（已删除 board_list.BOARD_CFG，完全不依赖）=====
from cable_list import CABLE_DATA
from flip_list import FLIP_LIST
from img_code_map import map_cable, map_board, map_conn, map_if, map_mode
from scene_list import SCENE_LIST
from machine_slot_map import MACHINE_CONFIG
from board_excel_map import BOARD_EXCEL_PATH_MAP


class PinConvertApp:
    def __init__(self, root):
        self.root = root
        self.root.title("线缆引脚定义生成工具v1.9")
        self.root.geometry("1080x520")

        # ========== 新增：顶部标签页容器 ==========
        self.notebook = ttk.Notebook(root)
        self.notebook.pack(fill="both", expand=True, padx=5, pady=5)

        # 标签页1：引脚生成页面【增加横竖滚动】
        self.tab1 = tk.Frame(self.notebook)
        self.notebook.add(self.tab1, text="引脚定义生成")

        tab1_canvas = tk.Canvas(self.tab1)
        scroll_v = ttk.Scrollbar(self.tab1, orient="vertical", command=tab1_canvas.yview)
        scroll_h = ttk.Scrollbar(self.tab1, orient="horizontal", command=tab1_canvas.xview)
        tab1_canvas.configure(yscrollcommand=scroll_v.set, xscrollcommand=scroll_h.set)

        tab1_canvas.grid(row=0, column=0, sticky="nsew")
        scroll_v.grid(row=0, column=1, sticky="ns")
        scroll_h.grid(row=1, column=0, sticky="ew")
        self.tab1.rowconfigure(0, weight=1)
        self.tab1.columnconfigure(0, weight=1)

        master_frame = tk.Frame(tab1_canvas)
        tab1_canvas.create_window((0, 0), window=master_frame, anchor="nw")

        def update_scroll_region(event):
            tab1_canvas.configure(scrollregion=tab1_canvas.bbox("all"))

        master_frame.bind("<Configure>", update_scroll_region)

        # 鼠标滚轮：仅在tab1画布上悬浮才滚动
        def _on_mousewheel(event):
            if tab1_canvas.winfo_containing(event.x_root, event.y_root) == tab1_canvas:
                tab1_canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        self.root.bind_all("<MouseWheel>", _on_mousewheel)

        master_frame.columnconfigure(0, weight=1)
        master_frame.columnconfigure(1, weight=0)

        self.scroll_text = "欢迎使用线缆引脚定义生成工具    该工具仅限御渡半导体内部人员使用    如有优化建议或遇使用问题请联系开发人员    "
        self.scroll_pos = 0

        self.cable_data = CABLE_DATA
        self.flip_list = FLIP_LIST
        self.board_excel_map = BOARD_EXCEL_PATH_MAP

        # 板卡Excel缓存 key=板卡名
        # 【优化后缓存结构】
        # {
        #   "__full_conn_list": [完整表头字符串1, 完整表头字符串2...],
        #   "C1(花面朝上)(DPS1_M)": {pin:net},
        #   "C1(花面朝下)(DPS1_M)": {pin:net}
        # }
        self._board_excel_cache = {}

        # 板卡下拉列表直接取自board_excel_map的keys，不再依赖board_list.py
        self.board_names = list(self.board_excel_map.keys())
        self.cable_list = list(self.cable_data.keys())
        self.line_widgets = []
        self.selected_row_idx = 0

        # ====第0行：线缆选择====
        row1 = tk.Frame(master_frame)
        row1.grid(row=0, column=0, columnspan=2, pady=20, padx=10, sticky="ew")
        tk.Label(row1, text="转接线缆选择：").pack(side="left")
        self.cable_combo = ttk.Combobox(row1, values=self.cable_list, state="readonly", width=20)
        self.cable_combo.pack(side="left", padx=8)
        self.cable_combo.bind("<<ComboboxSelected>>", self.on_cable_changed)
        if len(self.cable_list) > 0:
            self.cable_combo.current(0)

        tk.Label(row1, text="线缆长度：").pack(side="left", padx=(20, 0))
        self.cable_len_combo = ttk.Combobox(row1, state="readonly", width=10)
        self.cable_len_combo.pack(side="left", padx=8)
        self.cable_len_combo.bind("<<ComboboxSelected>>", self.on_length_changed)

        tk.Label(row1, text="线缆物料编码：").pack(side="left", padx=(20, 0))
        self.cable_code_var = tk.StringVar()
        tk.Label(row1, textvariable=self.cable_code_var, width=12, bg="white", relief="solid").pack(side="left", padx=8)

        self.cable_img_btn = ttk.Button(row1, text="查看线缆", command=self.show_cable_image)
        self.cable_img_btn.pack(side="left", padx=(10, 0))

        # 第1行：表格 + 教程
        table_area = tk.Frame(master_frame)
        table_area.grid(row=1, column=0, padx=10, sticky="nw")
        self.grid_container = tk.Frame(table_area)
        self.grid_container.pack(anchor="nw")

        tutorial_frame = tk.LabelFrame(master_frame, text="使用教程", padx=8, pady=0)
        tutorial_frame.grid(row=1, column=1, padx=(0, 10), sticky="n")

        tutorial_text = (
            "1. 选择对应的转接线缆\n"
            "2. 逐行选择和线缆对接的板卡型号和连接器位号(连接器自动从板卡Excel解析)\n"
            "3. 选择线缆的对接方式（花面朝上/朝下）\n"
            "4. 点击【查看板卡】查看当前行板卡示意图\n"
            "5. 点击【显示对接示意图】查看当前选中行的板卡连接器示意图\n"
            "6. 点击\"选择保存位置\"设置输出路径\n"
            "7. 点击\"生成引脚定义CSV文件\"\n"
            "8. CSV文件的G列为线缆输出端连接器PIN脚位号 \n"
            "9. CSV文件的C列为线缆输出端连接器PIN脚对应的网络名【网络全部读取自板卡Excel】 \n"
        )
        tk.Label(tutorial_frame, text=tutorial_text, justify="left",
                 font=("微软雅黑", 9), fg="#333333").pack(anchor="w")

        # 表格表头
        tk.Label(self.grid_container, text="板卡型号", width=16).grid(row=0, column=0, padx=8, sticky="w")
        tk.Label(self.grid_container, text="板卡连接器位号", width=28).grid(row=0, column=1, padx=8, sticky="w")
        tk.Label(self.grid_container, text="线缆连接器位号", width=12).grid(row=0, column=2, padx=8, sticky="w")
        tk.Label(self.grid_container, text="对接方式", width=12).grid(row=0, column=3, padx=8, sticky="w")
        tk.Label(self.grid_container, text="操作", width=10).grid(row=0, column=4, padx=8, sticky="w")

        img_btn_frame = tk.Frame(master_frame)
        img_btn_frame.grid(row=2, column=0, columnspan=1, pady=(0, 4), padx=10, sticky="w")
        self.global_img_btn = ttk.Button(img_btn_frame, text="显示对接示意图", command=self.show_selected_row_image)
        self.global_img_btn.pack(side="left")

        row_path = tk.Frame(master_frame)
        row_path.grid(row=3, column=0, columnspan=2, pady=(4, 8), padx=10, sticky="ew")
        tk.Label(row_path, text="输出文件路径：").pack(side="left")
        self.save_path_var = tk.StringVar()
        tk.Entry(row_path, textvariable=self.save_path_var, width=56).pack(side="left", padx=6)
        tk.Button(row_path, text="选择保存位置", command=self.select_save_path).pack(side="left")

        gen_btn = tk.Button(master_frame, text="生成转接后引脚定义CSV文件", command=self.generate_csv, bg="#337ab7",
                            fg="white", height=2)
        gen_btn.grid(row=4, column=0, columnspan=2, pady=20)

        status_bar = tk.Frame(master_frame, relief="sunken", bd=1, bg="#f0f0f0")
        status_bar.grid(row=5, column=0, columnspan=2, sticky="ew")
        tk.Label(status_bar, text="作者：冯志鹏", bg="#f0f0f0", font=("微软雅黑", 9)).pack(side="left", padx=10)
        self.scroll_label = tk.Label(status_bar, text="", bg="#f0f0f0", font=("微软雅黑", 9), fg="#333333")
        self.scroll_label.pack(side="left", fill="x", expand=True)
        tk.Label(status_bar, text="版本：v1.9", bg="#f0f0f0", font=("微软雅黑", 9)).pack(side="right", padx=10)

        self.update_scroll()
        self.on_cable_changed()

        # ========== 标签页2：接线装配说明==========
        self.tab2 = tk.Frame(self.notebook)
        self.notebook.add(self.tab2, text="接线装配说明(待开发)")
        tab2_main = tk.Frame(self.tab2)
        tab2_main.pack(fill="both", expand=True, padx=20, pady=20)

        scene_frame = tk.Frame(tab2_main)
        scene_frame.pack(anchor="w", pady=(0, 10))
        tk.Label(scene_frame, text="装配场景选择：", font=("微软雅黑", 11)).pack(side="left")
        self.scene_combo = ttk.Combobox(scene_frame, state="readonly", width=20)
        self.scene_combo.pack(side="left", padx=10)
        scene_btn = ttk.Button(scene_frame, text="查看装配示意图", command=self.show_scene_image)
        scene_btn.pack(side="left", padx=10)
        self.scene_combo.bind("<<ComboboxSelected>>", self.on_scene_selected)

        doc_frame = tk.LabelFrame(tab2_main, text="该场景关联文档", padx=10, pady=10)
        doc_frame.pack(fill="both", expand=True)
        list_container = tk.Frame(doc_frame)
        list_container.pack(fill="both", expand=True)
        self.scene_doc_listbox = tk.Listbox(list_container, font=("微软雅黑", 9))
        scrollbar = ttk.Scrollbar(list_container, orient="vertical", command=self.scene_doc_listbox.yview)
        self.scene_doc_listbox.config(yscrollcommand=scrollbar.set)
        self.scene_doc_listbox.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        self.scene_doc_listbox.bind("<Double-Button-1>", lambda e: self.open_selected_scene_doc())

        doc_btn_frame = tk.Frame(doc_frame)
        doc_btn_frame.pack(pady=(8, 0), anchor="w")
        self.open_doc_btn = ttk.Button(doc_btn_frame, text="打开选中文档", command=self.open_selected_scene_doc)
        self.open_doc_btn.pack()

        self.scene_cfg = SCENE_LIST
        scene_names = list(self.scene_cfg.keys())
        self.scene_combo["values"] = scene_names
        if scene_names:
            self.scene_combo.current(0)
            self.on_scene_selected()

        # ======================标签页3：软硬件通道映射======================
        self.tab3 = tk.Frame(self.notebook)
        self.notebook.add(self.tab3, text="软硬件通道映射")
        tab3_main = tk.Frame(self.tab3)
        tab3_main.pack(fill="both", expand=True, padx=30, pady=20)

        row_ch_top = tk.Frame(tab3_main)
        row_ch_top.pack(anchor="w", pady=(0, 8))
        tk.Label(row_ch_top, text="机台型号：", font=("微软雅黑", 11)).pack(side="left")
        self.machine_combo = ttk.Combobox(row_ch_top, state="readonly", width=18)
        self.machine_combo.pack(side="left", padx=(10, 20))
        self.machine_combo.bind("<<ComboboxSelected>>", self.on_machine_selected)

        tk.Label(row_ch_top, text="槽位号：", font=("微软雅黑", 11)).pack(side="left")
        self.slot_combo = ttk.Combobox(row_ch_top, state="readonly", width=15)
        self.slot_combo.pack(side="left", padx=(10, 20))
        self.slot_combo.bind("<<ComboboxSelected>>", self.on_slot_selected)

        tk.Label(row_ch_top, text="板卡名称：", font=("微软雅黑", 11)).pack(side="left")
        self.ch_board_var = tk.StringVar()
        tk.Label(row_ch_top, textvariable=self.ch_board_var, width=20,
                 bg="white", relief="solid", anchor="w", font=("微软雅黑", 11)).pack(side="left", padx=10)

        self.ch_form_frame = tk.LabelFrame(tab3_main, text="通道映射信息（对应Excel表格）", padx=12, pady=8)
        self.ch_form_frame.pack(fill="both", expand=True, pady=(4, 10), anchor="w")

        row_btn = tk.Frame(tab3_main)
        row_btn.pack(anchor="w", pady=(0, 0))
        self.open_ch_excel_btn = ttk.Button(row_btn, text="打开机台通道映射Excel", command=self.open_channel_map_file)
        self.open_ch_excel_btn.pack()

        self.ch_excel_rows = []
        self.ch_column_names = []
        self.ch_key_column = ""
        self.ch_form_rows = {}
        self.ch_key_combo = None
        self.current_excel_path = ""
        self._excel_cache = {}
        self._ch_req_id = 0

        self.machine_cfg = MACHINE_CONFIG
        machine_names = list(self.machine_cfg.keys())
        self.machine_combo["values"] = machine_names
        if machine_names:
            self.machine_combo.current(0)
            self.on_machine_selected()

    def _load_board_excel_cache(self, board_name):
        """【修复版】加载板卡Excel，解决表头空单元格导致连接器网络错位"""
        if board_name in self._board_excel_cache:
            return self._board_excel_cache[board_name]

        if board_name in self.board_excel_map:
            excel_path = resource_path(self.board_excel_map[board_name])
        else:
            excel_path = resource_path(f"conn_excel/{board_name}.xlsx")

        if not os.path.exists(excel_path):
            raise FileNotFoundError(f"板卡[{board_name}]对应的Excel文件不存在：{excel_path}")

        wb = load_workbook(excel_path, data_only=True)
        ws = wb["Sheet1"]
        all_rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not all_rows:
            raise Exception(f"{excel_path} Sheet1为空")

        header_row = all_rows[0]
        # ========== 改动点1：建立 {连接器全名:原始列索引} 映射 ==========
        conn_col_map = {}
        for col_idx, cell_text in enumerate(header_row):
            if cell_text is None:
                continue
            col_str = str(cell_text).strip()
            if not col_str:
                continue
            conn_col_map[col_str] = col_idx

        full_conn_list = list(conn_col_map.keys())
        board_conn_data = {
            "__full_conn_list": full_conn_list
        }
        # 初始化每个连接器的pin字典
        for conn_name in full_conn_list:
            board_conn_data[conn_name] = {}

        if len(full_conn_list) == 0:
            raise Exception(f"板卡[{board_name}]Excel表头未解析到任何连接器，请检查表头")

        # ========== 改动点2：按映射的原始列号取值，不再用循环col_idx偏移 ==========
        for row in all_rows[1:]:
            pin_sig = row[0]
            if pin_sig is None or str(pin_sig).strip() == "":
                continue
            pin_sig = str(pin_sig).strip()
            # 遍历每一个连接器，取它绑定的原始列号
            for conn_name, real_col_idx in conn_col_map.items():
                if real_col_idx >= len(row):
                    net_str = ""
                else:
                    net_val = row[real_col_idx]
                    net_str = str(net_val).strip() if net_val is not None else ""
                board_conn_data[conn_name][pin_sig] = net_str

        self._board_excel_cache[board_name] = board_conn_data
        return board_conn_data

    def _get_board_network(self, board_name, full_conn_text, lookup_pin):
        """【优化】仅3入参，只查表，不处理朝向、翻转"""
        try:
            board_data = self._load_board_excel_cache(board_name)
        except Exception as e:
            print(f"读取板卡Excel异常:{e}")
            return ""

        if full_conn_text not in board_data:
            return ""
        pin_net_dict = board_data[full_conn_text]
        net = pin_net_dict.get(lookup_pin, "")
        return net

    def on_board_selected(self, board_combo, conn_combo, row_idx):
        """【改】连接器下拉填充【完整表头文本】"""
        self.selected_row_idx = row_idx
        selected_board = board_combo.get()
        try:
            board_data = self._load_board_excel_cache(selected_board)
            conn_full_list = board_data.get("__full_conn_list", [])
        except Exception as e:
            err_msg = f"读取板卡【{selected_board}】Excel失败：\n{e}"
            messagebox.showerror("板卡Excel解析错误", err_msg)
            print(err_msg)
            conn_full_list = []

        conn_combo["values"] = conn_full_list
        if len(conn_full_list) > 0:
            conn_combo.current(0)
        else:
            conn_combo.set("")

    def on_machine_selected(self, event=None):
        machine_name = self.machine_combo.get()
        machine_info = self.machine_cfg.get(machine_name, {})
        slot_dict = machine_info.get("slots", {})
        slot_list = list(slot_dict.keys())
        self.slot_combo["values"] = slot_list
        if slot_list:
            self.slot_combo.current(0)
        else:
            self.slot_combo.set("")
        self.on_slot_selected()

    def on_slot_selected(self, event=None):
        self._reset_ch_state()
        machine_name = self.machine_combo.get()
        slot_name = self.slot_combo.get()
        machine_info = self.machine_cfg.get(machine_name, {})
        slot_dict = machine_info.get("slots", {})
        board_name = slot_dict.get(slot_name, "")
        self.ch_board_var.set(board_name)

        raw_excel = machine_info.get("excel_file", "").strip()
        sheet_name = slot_name.strip()
        board_columns = machine_info.get("board_columns", {})
        col_cfg = board_columns.get(board_name, {})
        columns = col_cfg.get("columns", []) or machine_info.get("columns", [])
        key_column = col_cfg.get("key_column", "").strip() or machine_info.get("key_column", "").strip()

        if not raw_excel:
            self._show_form_status(f"机台[{machine_name}]未配置 excel_file（Excel文件路径）")
            return
        if not columns:
            self._show_form_status(f"板卡[{board_name}]未配置 columns（请检查 board_columns）")
            return

        self.current_excel_path = resource_path(raw_excel)
        self.ch_column_names = [c for c in columns if str(c).strip()]
        self.ch_key_column = key_column

        if not self.ch_column_names:
            self._show_form_status(f"板卡[{board_name}]的columns配置为空或全是空格，请检查 board_columns")
            return

        if not os.path.exists(self.current_excel_path):
            is_unc = self.current_excel_path.startswith(r"\\")
            tip = "\n请确认已连接共享盘、有访问权限" if is_unc else ""
            self._show_form_status(f"找不到通道映射Excel：\n{self.current_excel_path}{tip}")
            return

        self._ch_req_id += 1
        req_id = self._ch_req_id
        cache = self._excel_cache.get(self.current_excel_path, {})
        if sheet_name in cache:
            self.ch_excel_rows = cache[sheet_name]
            self._build_ch_form()
            return

        self._show_form_status("正在读取Excel，请稍候...")

        def _worker():
            try:
                rows = self._read_excel_rows(self.current_excel_path, sheet_name)
                if rows:
                    c = self._excel_cache.get(self.current_excel_path, {})
                    c[sheet_name] = rows
                    self._excel_cache[self.current_excel_path] = c
                self.root.after(0, lambda: self._on_excel_loaded(rows, sheet_name, req_id))
            except Exception as e:
                err = str(e)
                self.root.after(0, lambda: self._show_form_status(f"读取Excel失败：\n{err}"))

        threading.Thread(target=_worker, daemon=True).start()

    def _on_excel_loaded(self, rows, sheet_name, req_id=None):
        if req_id is not None and req_id != self._ch_req_id:
            return
        self.ch_excel_rows = rows
        if not rows:
            self._show_form_status(
                f"sheet[{sheet_name}]中未找到可显示的数据行。\n"
                f"请检查：\n"
                f"1. 该sheet是否有数据\n"
                f"2. 配置的列名[{self.ch_column_names}]是否和Excel表头完全一致（注意空格/错别字）"
            )
            return
        self._build_ch_form()

    def _clear_ch_form(self):
        for w in self.ch_form_frame.winfo_children():
            w.destroy()
        self.ch_form_rows = {}
        self.ch_key_combo = None

    def _reset_ch_state(self):
        self._clear_ch_form()
        self.ch_excel_rows = []
        self.ch_column_names = []
        self.ch_key_column = ""

    def _show_form_status(self, msg):
        self._clear_ch_form()
        tk.Label(self.ch_form_frame, text=msg, fg="#a00",
                 font=("微软雅黑", 10), justify="left").grid(row=0, column=0, padx=6, pady=10, sticky="w")

    def _read_excel_rows(self, full_path, sheet_name=None):
        wb = load_workbook(filename=full_path, data_only=True)
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                wb.close()
                raise Exception(f"Excel中找不到sheet：{sheet_name}\n现有sheet：{wb.sheetnames}")
            ws = wb[sheet_name]
        else:
            ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not all_rows:
            raise Exception("Excel文件为空或没有读到任何行")

        header_row_idx = None
        header = None
        for i, row in enumerate(all_rows):
            cell_vals = set()
            for v in row:
                if v is not None and str(v).strip() != "":
                    cell_vals.add(str(v).strip())
            if all(col in cell_vals for col in self.ch_column_names):
                header_row_idx = i
                header = row
                break
        if header is None:
            preview = []
            for i, r in enumerate(all_rows[:3]):
                cells = [str(v).strip() for v in r if v is not None and str(v).strip() != ""]
                preview.append(f"第{i+1}行: {cells}")
            raise Exception(
                f"找不到包含所有配置列名的表头行。\n"
                f"需要的列：{self.ch_column_names}\n"
                f"Excel前3行内容：\n" + "\n".join(preview)
            )
        header_str = [str(v).strip() if v is not None else "" for v in header]
        col_index_map = {}
        missing_cols = []
        for col_name in self.ch_column_names:
            if col_name in header_str:
                col_index_map[col_name] = header_str.index(col_name)
            else:
                missing_cols.append(col_name)
        if missing_cols:
            raise Exception(
                f"Excel表头缺少配置列：{missing_cols}\n"
                f"实际表头：{header_str}"
            )
        rows = []
        for row in all_rows[header_row_idx + 1:]:
            rec = {}
            for col_name, idx in col_index_map.items():
                val = row[idx] if idx < len(row) else None
                rec[col_name] = str(val).strip() if (val is not None and str(val).strip() != "") else ""
            if any(rec.values()):
                rows.append(rec)
        return rows

    def _build_ch_form(self):
        if not self.ch_column_names:
            return
        self._clear_ch_form()
        for i, col in enumerate(self.ch_column_names):
            tk.Label(self.ch_form_frame, text=col, font=("微软雅黑", 11)).grid(row=i, column=0, sticky="w", padx=6, pady=6)
            if col == self.ch_key_column:
                values = list(dict.fromkeys(r.get(self.ch_key_column, "") for r in self.ch_excel_rows if r.get(self.ch_key_column, "")))
                combo = ttk.Combobox(self.ch_form_frame, state="readonly", width=30)
                combo["values"] = values
                combo.grid(row=i, column=1, sticky="w", padx=6, pady=6)
                combo.bind("<<ComboboxSelected>>", self.on_ch_selected)
                self.ch_key_combo = combo
                self.ch_form_rows[col] = combo
            else:
                var = tk.StringVar()
                lbl = tk.Label(self.ch_form_frame, textvariable=var, width=30, bg="white", relief="solid", anchor="w", font=("微软雅黑", 11))
                lbl.grid(row=i, column=1, sticky="w", padx=6, pady=6)
                self.ch_form_rows[col] = var
        if self.ch_key_combo is not None and self.ch_key_combo["values"]:
            self.ch_key_combo.current(0)
            self.on_ch_selected()

    def on_ch_selected(self, event=None):
        if self.ch_key_combo is None or not self.ch_key_column:
            return
        key_val = self.ch_key_combo.get()
        matched = None
        for r in self.ch_excel_rows:
            if str(r.get(self.ch_key_column, "")) == key_val:
                matched = r
                break
        for col, widget in self.ch_form_rows.items():
            if col == self.ch_key_column:
                continue
            val = matched.get(col, "") if matched else ""
            widget.set(str(val) if val not in (None, "nan") else "")

    def open_channel_map_file(self):
        full_path = self.current_excel_path
        if not full_path:
            messagebox.showinfo("提示", "当前机台未配置Excel文件路径")
            return
        is_unc = full_path.startswith(r"\\")
        if not os.path.exists(full_path):
            tip = "\n⚠️网络共享路径，请确认已连接共享盘、有访问权限" if is_unc else ""
            messagebox.showerror("文件不存在", f"找不到通道映射Excel：\n{full_path}{tip}")
            return
        try:
            os.startfile(full_path)
        except Exception as e:
            err_msg = str(e)
            if is_unc:
                err_msg += "\n可能原因：未映射共享盘/无访问权限/共享断开"
            messagebox.showerror("打开失败", f"无法打开Excel：{err_msg}")

    def on_scene_selected(self, event=None):
        scene_name = self.scene_combo.get()
        scene_info = self.scene_cfg.get(scene_name, {})
        doc_list = scene_info.get("docs", [])
        self.scene_doc_listbox.delete(0, tk.END)
        self._scene_docs_cache = doc_list
        for doc_item in doc_list:
            display_name = doc_item.get("name", "")
            self.scene_doc_listbox.insert(tk.END, display_name)

    def open_selected_scene_doc(self):
        sel_idx = self.scene_doc_listbox.curselection()
        if not sel_idx:
            messagebox.showinfo("提示", "请先在列表选中一个文档")
            return
        idx = sel_idx[0]
        doc_item = self._scene_docs_cache[idx]
        raw_path = doc_item.get("path", "").strip()
        full_path = resource_path(raw_path)
        is_unc = full_path.startswith(r"\\")
        if not os.path.exists(full_path):
            tip = "\n⚠️网络共享路径，请确认已连接共享盘、有访问权限" if is_unc else ""
            messagebox.showerror("文件不存在", f"找不到文档：\n{full_path}{tip}")
            return
        try:
            os.startfile(full_path)
        except Exception as e:
            err_msg = str(e)
            if is_unc:
                err_msg += "\n可能原因：未映射共享盘/无访问权限/共享断开"
            messagebox.showerror("打开失败", f"无法打开文档：{err_msg}")

    def show_cable_image(self):
        cable_name = self.cable_combo.get()
        base_name = f"cable_{cable_name}"
        img_folder = resource_path("conn_img")
        png_path = os.path.join(img_folder, f"{base_name}.png")
        gif_path = os.path.join(img_folder, f"{base_name}.gif")
        self.show_full_conn_image(base_name, png_path, gif_path)

    def show_board_image(self, board_name):
        base_name = f"board_{board_name}"
        img_folder = resource_path("conn_img")
        png_path = os.path.join(img_folder, f"{base_name}.png")
        gif_path = os.path.join(img_folder, f"{base_name}.gif")
        self.show_full_conn_image(base_name, png_path, gif_path)

    def show_scene_image(self):
        selected_scene = self.scene_combo.get()
        base_name = f"scene_{selected_scene}"
        img_folder = resource_path("conn_img")
        png_path = os.path.join(img_folder, f"{base_name}.png")
        gif_path = os.path.join(img_folder, f"{base_name}.gif")
        self.show_full_conn_image(base_name, png_path, gif_path)

    def on_length_changed(self, event=None):
        cable_name = self.cable_combo.get()
        cfg = self.cable_data.get(cable_name, {})
        len_dict = cfg.get("length_list", {})
        selected_len = self.cable_len_combo.get()
        code = len_dict.get(selected_len, "")
        self.cable_code_var.set(code)

    def on_mode_changed(self, row_idx, event):
        self.selected_row_idx = row_idx

    def on_cable_changed(self, event=None):
        cable_name = self.cable_combo.get()
        if cable_name not in self.cable_data:
            return
        cfg = self.cable_data[cable_name]
        input_list = cfg["input_if"]

        len_dict = cfg.get("length_list", {})
        len_options = list(len_dict.keys())
        default_len = cfg.get("default_len", "")
        self.cable_len_combo["values"] = len_options
        if len_options:
            self.cable_len_combo.set(default_len)
        else:
            self.cable_len_combo.set("")
        self.on_length_changed()

        for w in self.grid_container.winfo_children():
            w.destroy()
        self.line_widgets.clear()
        self.selected_row_idx = 0

        tk.Label(self.grid_container, text="板卡型号", width=16).grid(row=0, column=0, padx=8, sticky="w")
        tk.Label(self.grid_container, text="板卡连接器位号", width=28).grid(row=0, column=1, padx=8, sticky="w")
        tk.Label(self.grid_container, text="线缆连接器位号", width=12).grid(row=0, column=2, padx=8, sticky="w")
        tk.Label(self.grid_container, text="对接方式", width=12).grid(row=0, column=3, padx=8, sticky="w")
        tk.Label(self.grid_container, text="操作", width=10).grid(row=0, column=4, padx=8, sticky="w")

        for row_idx, item in enumerate(input_list, start=1):
            cb_board = ttk.Combobox(self.grid_container, values=self.board_names, state="readonly", width=16)
            if len(self.board_names) > 0:
                cb_board.current(0)
            cb_board.grid(row=row_idx, column=0, padx=8, pady=3, sticky="w")

            cb_conn = ttk.Combobox(self.grid_container, values=[], state="readonly", width=36)
            cb_conn.grid(row=row_idx, column=1, padx=8, pady=3, sticky="w")

            lbl_if = tk.Label(self.grid_container, text=item, width=12)
            lbl_if.grid(row=row_idx, column=2, padx=8, pady=3, sticky="w")

            cb_mode = ttk.Combobox(self.grid_container, values=["花面朝上", "花面朝下"], state="readonly", width=12)
            cb_mode.current(0)
            cb_mode.grid(row=row_idx, column=3, padx=8, pady=3, sticky="w")

            btn_board_img = ttk.Button(self.grid_container, text="查看板卡", width=8,
                                       command=lambda b=cb_board: self.show_board_image(b.get()))
            btn_board_img.grid(row=row_idx, column=4, padx=8, pady=3)

            cb_board.bind("<<ComboboxSelected>>",
                          lambda e, b=cb_board, c=cb_conn, r=row_idx: self.on_board_selected(b, c, r))
            cb_mode.bind("<<ComboboxSelected>>", lambda e, r=row_idx: self.on_mode_changed(r, e))

            self.line_widgets.append({
                "if_name": item,
                "board_combo": cb_board,
                "conn_combo": cb_conn,
                "mode_combo": cb_mode,
            })
            self.on_board_selected(cb_board, cb_conn, row_idx)

    def show_selected_row_image(self):
        if len(self.line_widgets) == 0:
            messagebox.showwarning("提示", "当前线缆没有输入接口！")
            return
        cable_name = self.cable_combo.get()
        cable_id = map_cable.get(cable_name, 99)
        parts_num = [str(cable_id)]
        for line in self.line_widgets:
            board_name = line["board_combo"].get()
            full_conn_text = line["conn_combo"].get()
            wire_if = line["if_name"]
            mode = line["mode_combo"].get()
            b_id = map_board.get(board_name, 99)
            # 从完整文本提取sig用于map_conn
            m = re.match(r"([A-Z]+\d+)\((花面朝上|花面朝下)\)", full_conn_text)
            if m:
                conn_sig = m.group(1)
            else:
                conn_sig = full_conn_text
            c_id = 99
            for key, val in map_conn.items():
                if key in conn_sig:
                    c_id = val
                    break
            i_id = map_if.get(wire_if, 99)
            m_id = map_mode.get(mode, 99)
            seg_num = f"{b_id}{c_id}{i_id}{m_id}"
            parts_num.append(seg_num)
        base_name = "_".join(parts_num)
        img_folder = resource_path("conn_img")
        png_path = os.path.join(img_folder, f"{base_name}.png")
        gif_path = os.path.join(img_folder, f"{base_name}.gif")
        self.show_full_conn_image(base_name, png_path, gif_path)

    def show_full_conn_image(self, base_name, png_path, gif_path):
        img_win = tk.Toplevel(self.root)
        img_win.title(f"示意图 - {base_name}")
        img_win.after_id = None
        if os.path.exists(gif_path):
            img_win.frames = []
            img_win.frame_index = 0
            try:
                frame_num = 0
                while True:
                    frame = tk.PhotoImage(file=gif_path, format=f"gif -index {frame_num}")
                    img_win.frames.append(frame)
                    frame_num += 1
            except tk.TclError:
                pass
            if len(img_win.frames) == 0:
                messagebox.showerror("GIF加载失败", "GIF无有效帧！")
                img_win.destroy()
                return
            label_img = tk.Label(img_win)
            label_img.pack(expand=True)

            def animate():
                if not img_win.winfo_exists():
                    return
                frame = img_win.frames[img_win.frame_index]
                label_img.config(image=frame)
                img_win.frame_index += 1
                if img_win.frame_index >= len(img_win.frames):
                    img_win.frame_index = 0
                img_win.after_id = img_win.after(150, animate)

            animate()
        elif os.path.exists(png_path):
            try:
                photo = tk.PhotoImage(file=png_path)
                img_win.photo = photo
                label_img = tk.Label(img_win, image=photo)
                label_img.pack(expand=True)
            except Exception as e:
                messagebox.showerror("PNG加载失败", f"{e}")
                img_win.destroy()
                return
        else:
            messagebox.showerror("图片不存在",
                                 f"找不到示意图文件：\n{base_name}.gif 或 {base_name}.png\n请检查接线方式是否正确")
            img_win.destroy()
            return

        def on_close():
            if img_win.after_id is not None:
                img_win.after_cancel(img_win.after_id)
            img_win.destroy()

        img_win.protocol("WM_DELETE_WINDOW", on_close)

    def select_save_path(self):
        path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV文件", "*.csv"), ("所有文件", "*.*")],
            title="保存引脚定义文件"
        )
        if path:
            self.save_path_var.set(path)

    def generate_csv(self):
        cable_name = self.cable_combo.get()
        cable_length = self.cable_len_combo.get()
        cable_code = self.cable_code_var.get()
        out_path = self.save_path_var.get().strip()
        if not out_path:
            messagebox.showwarning("提示", "请先选择输出路径！")
            return

        cfg = self.cable_data[cable_name]
        pin_map = cfg["pin_mapping"]
        pin_list = list(pin_map.items())

        line_setting = {}
        for line in self.line_widgets:
            line_setting[line["if_name"]] = {
                "board": line["board_combo"].get(),
                "connector": line["conn_combo"].get(),
                "mode": line["mode_combo"].get()
            }

        headers = ["板卡型号", "板卡连接器位号", "网络", "线缆连接器位号", "输入PIN脚", "对接方式", "输出PIN脚", "线缆型号",
                   "线缆长度", "线缆物料编码"]
        rows = []

        for dest_pin, src_pin in pin_list:
            if dest_pin == "F1_E38":
                aaa=0
            if not src_pin:
                if_name = ""
                pin_num = ""
            elif "_" in src_pin:
                idx = src_pin.rfind("_")
                if_name = src_pin[:idx]
                pin_num = src_pin[idx + 1:]
            else:
                if_name = ""
                pin_num = ""

            s = line_setting.get(if_name, {"board": "", "connector": "", "mode": ""})
            lookup_pin = pin_num
            conn_name = s["connector"]
            gui_mode = s["mode"]
            board_name = s["board"]

            # 上层在这里做翻转判断，查表函数不再处理朝向

            if conn_name and gui_mode and (gui_mode not in conn_name):
                flip_dict = self.flip_list.get("CONNECTOR_C", {})
                lookup_pin = flip_dict.get(lookup_pin, lookup_pin)

            # 【3参数调用】
            net_name = self._get_board_network(board_name, conn_name, lookup_pin)

            row = [s["board"], conn_name, net_name, if_name, src_pin, s["mode"], dest_pin, cable_name, cable_length,
                   cable_code]
            rows.append(row)

        try:
            with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                writer.writerows(rows)
            messagebox.showinfo("成功", "CSV文件生成完成")
        except Exception as e:
            messagebox.showerror("保存失败", str(e))

    def update_scroll(self):
        self.scroll_label.config(text=self.scroll_text[self.scroll_pos:] + self.scroll_text[:self.scroll_pos])
        self.scroll_pos += 1
        if self.scroll_pos >= len(self.scroll_text):
            self.scroll_pos = 0
        self.root.after(200, self.update_scroll)


if __name__ == "__main__":
    root = tk.Tk()
    app = PinConvertApp(root)
    root.mainloop()
